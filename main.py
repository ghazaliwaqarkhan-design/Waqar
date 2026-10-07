import glob
import hashlib
import io
import os
import random
import re
import time

import numpy as np
import streamlit as st
from google import genai
from google.genai import types

# =========================================================
# CONFIGURATION
# =========================================================
DATA_DIR = "data"
EMBED_MODEL = "gemini-embedding-001"  # runs on Google servers: no torch, light and fast
EMBED_DIM = 768
LLM_MODEL = "gemini-3.8-flash"  # preferred model; others are tried if it is busy
TOP_K = 3
IMAGE_DIR = "images"  # professor photos, faculty / campus photos
DOC_EXTS = (".txt", ".docx", ".pdf")  # files the chatbot can read from the data folder
IMAGE_EXTS = (".jpg", ".jpeg", ".png", ".webp")
IMAGE_CANDIDATES = 2  # how many photos are offered to the AI for each question
PHOTO_WIDTH = 240
MAX_CHUNK = 1500  # longest piece of text kept as one searchable chunk
MAX_QUESTIONS = 30  # per visit: protects the shared free quota when published

SYSTEM_PROMPT = """You are a warm, friendly student-support assistant for Nangarhar University. \
You chat like ChatGPT: natural, helpful and human, like a kind older classmate.

Rules:
1. Always reply in Pashto (Arabic script), in a natural, friendly, conversational tone.
2. Use emojis sometimes: about one per reply, and none when the topic is serious or the student is upset. \
Never put an emoji in every sentence.
3. For greetings, thanks and small talk, reply naturally and warmly, then offer to help with university questions.
4. For university questions, use ONLY the "Context from university documents" you are given. \
If the answer is not there, say honestly in Pashto that you do not have that information and suggest \
checking nu.edu.af or asking the university office. Never invent facts, numbers, dates or fees.
5. Use the earlier messages in the conversation to understand follow-up questions.
6. Keep answers short and clear. Use a short list only when listing several items.
7. If asked about something unrelated to the university (medical, legal, etc.), politely say you mainly \
help with Nangarhar University questions.
8. If anyone asks who made, created, built or designed this chatbot, answer in Pashto that it was designed by \
Web Designer Saeed ul Haq Chardiwal (write this name exactly, in English letters). \
If asked which AI technology you use, you may say it uses a Google Gemini language model.
9. Do NOT end your reply with an offer like "ask me more questions": the app adds that line by itself.
10. Photos: only show a photo that appears in the "Photos you may show" list of the message, and only when the \
student asks about that person or place (for example "who is the dean?" or "show me the library"). \
Never say a photo shows something that its caption does not say. If the student asks to see a photo and none \
is listed, say politely in Pashto that you do not have that photo yet."""

CLOSINGS = [
    "زه دلته یم چې ستاسو پوښتنو ته ځواب ووایم. که نورې پوښتنې لرئ، په آزاده توګه یې وکړئ 😊",
    "که کومه بله پوښتنه لرئ، په خوښۍ سره ځواب درکوم. په آزاده توګه وپوښتئ! 🌟",
    "زه همدلته یم! که د پوهنتون په اړه نورې پوښتنې لرئ، راته ولیکئ 🎓",
    "که مې ځواب ګټور و، نوره پوښتنه هم کولی شئ 🙌",
]

HISTORY_MESSAGES = 8  # how many previous messages the bot remembers

EXAMPLES = [
    "ننګرهار پوهنتون چېرته دی؟",
    "د کمپیوټر ساینس پوهنځی څه لري؟",
    "څنګه پوهنتون ته داخل شم؟",
    "د پوهنتون رسمي ویب پاڼه کومه ده؟",
]

# =========================================================
# PAGE CONFIG
# =========================================================
st.set_page_config(
    page_title="Nangarhar University AI",
    page_icon="🎓",
    layout="wide",
    initial_sidebar_state="expanded",
)

# =========================================================
# CSS
# =========================================================
CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Noto+Naskh+Arabic:wght@400;500;600;700&display=swap');

/* GLOBAL */
html, body, [class*="css"] {font-family: "Noto Naskh Arabic", "Segoe UI", sans-serif;}
[data-testid="stAppViewContainer"] {
    background:
        radial-gradient(circle at 10% 10%, rgba(20, 115, 90, 0.08), transparent 30%),
        radial-gradient(circle at 90% 10%, rgba(15, 76, 129, 0.08), transparent 30%),
        #f5f8fc;
}
.block-container {max-width: 1180px; padding-top: 1.5rem; padding-bottom: 5rem;}

/* HIDE STREAMLIT ELEMENTS */
#MainMenu, footer, [data-testid="stToolbar"], [data-testid="stDecoration"] {visibility: hidden;}

/* SIDEBAR */
[data-testid="stSidebar"] {
    background: linear-gradient(180deg, #ffffff 0%, #f7fafc 100%);
    border-right: 1px solid #e4eaf1;
}
[data-testid="stSidebar"] > div:first-child {padding-top: 1.2rem;}
.sidebar-brand {text-align: center; padding: 10px 5px 20px;}
.sidebar-logo {
    width: 70px; height: 70px; margin: auto; display: flex;
    align-items: center; justify-content: center; border-radius: 22px;
    background: linear-gradient(135deg, #0f4c81, #147b6b);
    color: white; font-size: 34px; box-shadow: 0 10px 25px rgba(15, 76, 129, 0.20);
}
.sidebar-title {margin-top: 12px; font-size: 1.25rem; font-weight: 700; color: #12344d;}
.sidebar-subtitle {color: #718096; font-size: 0.85rem;}

/* SIDEBAR BUTTONS */
[data-testid="stSidebar"] .stButton button {
    width: 100%; border-radius: 12px; border: 1px solid #dce4ed;
    background: white; color: #243b53; transition: 0.2s ease;
    font-family: "Noto Naskh Arabic", "Segoe UI", sans-serif;
}
[data-testid="stSidebar"] .stButton button:hover {
    border-color: #147b6b; color: #147b6b; transform: translateY(-1px);
    box-shadow: 0 5px 15px rgba(20, 123, 107, 0.10);
}

/* HERO */
.hero {
    position: relative; overflow: hidden;
    background: linear-gradient(135deg, #0b416f 0%, #126b6a 52%, #19846d 100%);
    color: white; border-radius: 26px; padding: 34px 30px; margin-bottom: 22px;
    box-shadow: 0 18px 40px rgba(15, 76, 129, 0.20);
}
.hero:before {
    content: ""; position: absolute; width: 230px; height: 230px; border-radius: 50%;
    right: -70px; top: -100px; background: rgba(255, 255, 255, 0.08);
}
.hero:after {
    content: ""; position: absolute; width: 160px; height: 160px; border-radius: 50%;
    left: -60px; bottom: -90px; background: rgba(255, 255, 255, 0.06);
}
.hero-content {position: relative; z-index: 2;}
.hero-icon {font-size: 42px; margin-bottom: 5px;}
.hero-title {direction: rtl; font-size: 2rem; font-weight: 700; line-height: 1.5;}
.hero-subtitle {font-family: "Segoe UI", sans-serif; font-size: 0.95rem; opacity: 0.88; margin-top: 5px;}
.status {
    display: inline-flex; align-items: center; margin-top: 18px; padding: 7px 13px;
    border-radius: 999px; background: rgba(255, 255, 255, 0.13); font-size: 0.85rem;
}
.status-dot {width: 8px; height: 8px; background: #4ade80; border-radius: 50%; margin-right: 7px;}

/* WELCOME CARD */
.welcome {
    direction: rtl; text-align: center; background: white; border: 1px solid #e5ebf2;
    border-radius: 20px; padding: 28px; margin: 10px 0 20px;
    box-shadow: 0 5px 18px rgba(20, 35, 50, 0.04);
}
.welcome-icon {font-size: 38px;}
.welcome-title {font-size: 1.3rem; font-weight: 700; color: #12344d; margin-top: 5px;}
.welcome-text {color: #697586; font-size: 0.96rem; margin-top: 7px; line-height: 2;}

/* API BOX */
.api-box {
    direction: rtl; text-align: right; background: white; border: 1px solid #e1e8f0;
    border-radius: 18px; padding: 20px; margin-bottom: 20px;
    box-shadow: 0 5px 20px rgba(20, 35, 50, 0.05);
}
.api-title {color: #12344d; font-size: 1.05rem; font-weight: 700; margin-bottom: 5px;}
.api-description {color: #718096; font-size: 0.88rem; line-height: 1.8; margin-bottom: 10px;}

/* CHAT */
[data-testid="stChatMessage"] {
    border-radius: 18px !important; border: 1px solid #e4eaf1 !important;
    padding: 12px 16px !important; margin-bottom: 12px !important;
    box-shadow: 0 4px 14px rgba(20, 35, 50, 0.035); background: white;
}
[data-testid="stChatMessage"] p, [data-testid="stChatMessage"] li {
    font-family: "Noto Naskh Arabic", "Segoe UI", sans-serif;
    font-size: 1.08rem; line-height: 2; direction: rtl; text-align: right; color: #17212b;
}

/* CHAT INPUT */
[data-testid="stChatInput"] {margin-top: 15px;}
[data-testid="stChatInput"] textarea {
    direction: rtl; text-align: right;
    font-family: "Noto Naskh Arabic", "Segoe UI", sans-serif; font-size: 1rem;
    border-radius: 16px !important; border: 1px solid #ccd6e2 !important;
    background: white !important; padding: 13px 16px !important;
}
[data-testid="stChatInput"] textarea:focus {
    border-color: #147b6b !important;
    box-shadow: 0 0 0 3px rgba(20, 123, 107, 0.10) !important;
}

/* EXAMPLE TITLE + FOOTER */
.example-title {
    direction: rtl; text-align: right; font-size: 1rem; font-weight: 700;
    color: #12344d; margin-top: 20px; margin-bottom: 8px;
}
.footer {
    direction: rtl; text-align: center; color: #8a96a3; font-size: 0.82rem;
    margin-top: 28px; padding: 15px; border-top: 1px solid #e5ebf2;
}

/* TYPING INDICATOR (shown while the answer is being prepared) */
.typing {display: flex; gap: 6px; padding: 10px 4px; direction: ltr;}
.typing span {
    width: 9px; height: 9px; border-radius: 50%; background: #147b6b; opacity: 0.3;
    animation: typing-bounce 1.2s infinite ease-in-out;
}
.typing span:nth-child(2) {animation-delay: 0.2s;}
.typing span:nth-child(3) {animation-delay: 0.4s;}
@keyframes typing-bounce {
    0%, 80%, 100% {opacity: 0.25; transform: translateY(0);}
    40% {opacity: 1; transform: translateY(-5px);}
}

/* RESPONSIVE */
@media (max-width: 768px) {
    .block-container {padding-left: 0.8rem; padding-right: 0.8rem;}
    .hero {padding: 25px 18px; border-radius: 20px;}
    .hero-title {font-size: 1.45rem;}
    .hero-subtitle {font-size: 0.8rem;}
    [data-testid="stChatMessage"] {border-radius: 15px !important;}
    [data-testid="stChatMessage"] p {font-size: 1rem;}
}
</style>
"""
st.markdown(CSS, unsafe_allow_html=True)

# =========================================================
# SESSION STATE
# =========================================================
if "history" not in st.session_state:
    st.session_state.history = []
if "api_key" not in st.session_state:
    st.session_state.api_key = ""
if "count" not in st.session_state:
    st.session_state.count = 0


# =========================================================
# SEARCH IN THE UNIVERSITY DATA (Gemini embeddings)
# =========================================================
def embed(api_key, texts, task):
    """Turn texts into normalized vectors using the Gemini embedding model."""
    client = get_client(api_key)
    cfg = types.EmbedContentConfig(task_type=task, output_dimensionality=EMBED_DIM)

    def call(contents):
        for attempt in range(3):
            try:
                return client.models.embed_content(model=EMBED_MODEL, contents=contents, config=cfg)
            except Exception as e:
                if attempt == 2 or not any(c in str(e) for c in ("503", "429", "UNAVAILABLE", "RESOURCE_EXHAUSTED")):
                    raise
                time.sleep(3)

    vectors = []
    for i in range(0, len(texts), 50):
        batch = texts[i : i + 50]
        try:
            vectors += [e.values for e in call(batch).embeddings]
        except Exception:
            for t in batch:  # fallback: one text per request
                vectors.append(call(t).embeddings[0].values)
    arr = np.array(vectors, dtype="float32")
    return arr / np.linalg.norm(arr, axis=1, keepdims=True)


def split_text(text):
    """Cut text into paragraph-sized chunks. Very long paragraphs are cut by lines."""
    chunks = []
    for para in re.split(r"\n\s*\n", text):
        para = para.strip()
        if len(para) <= 20:
            continue
        if len(para) <= MAX_CHUNK:
            chunks.append(para)
            continue
        lines = []
        for line in para.split("\n"):
            line = line.strip()
            while len(line) > MAX_CHUNK:  # one huge line: cut it
                lines.append(line[:MAX_CHUNK])
                line = line[MAX_CHUNK:]
            if line:
                lines.append(line)
        title = lines[0] if len(lines[0]) < 100 else ""
        piece = ""
        for line in lines:
            if piece and len(piece) + len(line) > MAX_CHUNK:
                chunks.append(piece)
                piece = (title + "\n" + line) if title and line != title else line
            else:
                piece = (piece + "\n" + line) if piece else line
        if piece:
            chunks.append(piece)
    return chunks


def merge_small(chunks, target=500):
    """Join neighbouring tiny pieces (PDF lines, headings) so each chunk carries enough meaning."""
    merged = []
    for c in chunks:
        if merged and len(merged[-1]) < target and len(merged[-1]) + len(c) <= MAX_CHUNK:
            merged[-1] += "\n" + c
        else:
            merged.append(c)
    return merged


def decode_text(raw):
    """Decode a text file saved by any Windows editor."""
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):  # UTF-16 (old Notepad "Unicode")
        return raw.decode("utf-16")
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1256", errors="ignore")  # Arabic-script "ANSI" files


def read_file(path):
    """Return the text inside a .txt, .docx or .pdf file (checks the real file type, not the name)."""
    with open(path, "rb") as f:
        raw = f.read()
    ext = os.path.splitext(path)[1].lower()
    if raw[:4] == b"PK\x03\x04" or ext == ".docx":
        kind = "docx"  # a Word file, even if somebody renamed it to .txt
    elif raw[:5] == b"%PDF-" or ext == ".pdf":
        kind = "pdf"
    else:
        kind = "txt"

    if kind == "docx":
        from docx import Document  # pip install python-docx

        doc = Document(io.BytesIO(raw))
        sections, current = [], []
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue  # blank lines do not split sections
            style = (p.style.name or "").lower() if p.style is not None else ""
            is_heading = style.startswith("heading") or style == "title" or (
                len(text) < 80 and re.match(r"^\d+[.)]\s+\S", text)
            )
            if is_heading and current:
                sections.append("\n".join(current))
                current = []
            current.append(text)
        for table in doc.tables:
            for row in table.rows:
                current.append(" | ".join(c.text.strip() for c in row.cells))
        if current:
            sections.append("\n".join(current))
        return "\n\n".join(sections)  # one chunk per section (heading + its content)
    if kind == "pdf":
        from pypdf import PdfReader  # pip install pypdf

        return "\n\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(raw)).pages)
    return decode_text(raw)


def load_images():
    """Find photos in the images folder. A caption comes from a .txt file with the same name."""
    items = []
    for path in sorted(glob.glob(os.path.join(IMAGE_DIR, "*"))):
        stem, ext = os.path.splitext(path)
        if ext.lower() not in IMAGE_EXTS:
            continue
        name = os.path.basename(path)
        caption = ""
        try:
            with open(stem + ".txt", "rb") as f:
                caption = decode_text(f.read()).strip()
        except Exception:
            pass
        if not caption:  # no caption file: use the file name, e.g. dr_ahmad_dean_cs.jpg
            caption = os.path.splitext(name)[0].replace("_", " ").replace("-", " ")
        items.append({"name": name, "path": path, "caption": caption})
    return items


def data_signature():
    """Changes whenever a file in data/ or images/ is added, edited or removed."""
    sig = []
    for folder in (DATA_DIR, IMAGE_DIR):
        for path in sorted(glob.glob(os.path.join(folder, "*"))):
            try:
                info = os.stat(path)
                sig.append((path, info.st_size, int(info.st_mtime)))
            except OSError:
                pass
    return tuple(sig)


@st.cache_resource
def load_index(api_key, signature):  # a new signature = the files changed = rebuild the index
    """Read documents + photo captions, embed them (saved on disk after the first time)."""
    chunks = []
    report = {"read": [], "skipped": [], "photos": []}
    for path in sorted(glob.glob(os.path.join(DATA_DIR, "*"))):
        name = os.path.basename(path)
        ext = os.path.splitext(name)[1].lower()
        if os.path.isdir(path) or name.startswith("~$") or name.lower() == "requirements.txt":
            continue
        if ext in IMAGE_EXTS:
            report["skipped"].append((name, "this is a photo: move it to the images folder"))
            continue
        if ext not in DOC_EXTS:
            report["skipped"].append((name, "file type not supported"))
            continue
        try:
            pieces = split_text(read_file(path))
            if ext == ".pdf":
                pieces = merge_small(pieces)
            if not pieces:
                report["skipped"].append((name, "no readable text inside (scanned PDF or empty)"))
                continue
            chunks += pieces
            report["read"].append((name, len(pieces)))
        except Exception as e:
            report["skipped"].append((name, f"could not be read: {e}"))

    images = load_images()
    report["photos"] = [img["name"] for img in images]
    texts = chunks + [img["caption"] for img in images]
    if not texts:
        return [], [], np.array([]), report

    key = hashlib.sha1((EMBED_MODEL + str(EMBED_DIM) + "\n".join(texts)).encode("utf-8")).hexdigest()
    cache_file = os.path.join(".cache", f"emb_{key}.npy")
    try:
        vectors = np.load(cache_file)
        if len(vectors) == len(texts):
            return chunks, images, vectors, report
    except Exception:
        pass

    vectors = embed(api_key, texts, "RETRIEVAL_DOCUMENT")
    try:
        os.makedirs(".cache", exist_ok=True)
        np.save(cache_file, vectors)
    except Exception:
        pass
    return chunks, images, vectors, report


def retrieve(question, api_key):
    """Return (best text chunks, best photo candidates) for the question."""
    chunks, images, vectors, _report = load_index(api_key, data_signature())
    if not chunks:
        return ["د پوهنتون د معلوماتو فایلونه پیدا نه شول."], []
    q = embed(api_key, [question], "RETRIEVAL_QUERY")[0]
    scores = vectors @ q
    n = len(chunks)
    top = np.argsort(scores[:n])[::-1][:TOP_K]
    pics = []
    if images:
        top_img = np.argsort(scores[n:])[::-1][:IMAGE_CANDIDATES]
        pics = [dict(images[i], id=f"P{k + 1}") for k, i in enumerate(top_img)]
    return [chunks[i] for i in top], pics


PHOTO_TAG = re.compile(r"\[\[photo:\s*([^\]]+?)\s*\]\]")


def photo_block(pics):
    """Tell the AI which photos exist, so it can decide whether to show one."""
    if not pics:
        return ""
    lines = "\n".join(f"- {p['id']}: {p['caption']}" for p in pics)
    return (
        "\n\nPhotos you may show (only if the student asks about that person or place; to show one, "
        "write [[photo:ID]] on its own line at the end of your reply, for example [[photo:P1]]):\n"
        + lines
    )


def split_photos(text, candidates):
    """Remove [[photo:...]] tags from the answer; return (clean text, photos to display)."""
    wanted = {m.strip().lower() for m in PHOTO_TAG.findall(text)}
    clean = PHOTO_TAG.sub("", text)
    clean = re.sub(r"\[\[[^\]]*$", "", clean).strip()  # unfinished tag while streaming
    shown = [c for c in candidates if c["id"].lower() in wanted]
    return clean, shown


def show_photos(photos):
    for p in photos:
        caption = p["caption"].split("\n")[0][:120]
        try:
            from PIL import Image, ImageOps

            img = ImageOps.exif_transpose(Image.open(p["path"]))  # fix phone-photo rotation
            img.thumbnail((700, 700))  # small file = fast page for students
            st.image(img, caption=caption, width=PHOTO_WIDTH)
        except Exception:
            st.image(p["path"], caption=caption, width=PHOTO_WIDTH)


# =========================================================
# GEMINI
# =========================================================
@st.cache_resource
def get_client(api_key):
    return genai.Client(api_key=api_key)


@st.cache_resource
def get_models(api_key):
    """Preferred model first, then other Flash models. Listed once, then remembered."""
    names = [LLM_MODEL]
    try:
        skip = ("image", "tts", "live", "audio", "embedding", "robotics", "computer")
        found = []
        for m in get_client(api_key).models.list():
            name = m.name.replace("models/", "")
            actions = getattr(m, "supported_actions", None) or ["generateContent"]
            if (
                "flash" in name
                and "generateContent" in actions
                and not any(s in name for s in skip)
                and name not in names
            ):
                found.append(name)
        names += sorted(found, reverse=True)
    except Exception:
        pass
    return names[:5]


def build_contents(history, question, context):
    """Previous messages first, then the new question together with the university context."""
    contents = []
    for item in history[-HISTORY_MESSAGES:]:
        contents.append(
            types.Content(
                role="user" if item["role"] == "user" else "model",
                parts=[types.Part(text=item["text"])],
            )
        )
    contents.append(
        types.Content(
            role="user",
            parts=[
                types.Part(
                    text=f"Context from university documents:\n{context}\n\nStudent message: {question}"
                )
            ],
        )
    )
    return contents


def stream_answer(question, context, api_key, history):
    """Yield the answer piece by piece so text appears while it is being written."""
    client = get_client(api_key)
    models = get_models(api_key)
    contents = build_contents(history, question, context)
    last_error = None
    for model in list(models):
        for attempt in range(2):
            started = False
            try:
                for chunk in client.models.generate_content_stream(
                    model=model,
                    contents=contents,
                    config=types.GenerateContentConfig(
                        system_instruction=SYSTEM_PROMPT, temperature=0.6
                    ),
                ):
                    if chunk.text:
                        started = True
                        yield chunk.text
                models.remove(model)  # remember the model that worked: use it first
                models.insert(0, model)
                return
            except Exception as e:
                if started:
                    yield "\n\n..."
                    return
                last_error = e
                if any(c in str(e) for c in ("503", "429", "UNAVAILABLE", "RESOURCE_EXHAUSTED")):
                    time.sleep(2)  # busy: short wait, retry the same model
                    continue
                break  # e.g. model not found: try the next model
    yield f"Error: {last_error}"


# =========================================================
# SIDEBAR
# =========================================================
with st.sidebar:
    st.html(
        """
        <div class="sidebar-brand">
            <div class="sidebar-logo">🎓</div>
            <div class="sidebar-title">ننګرهار پوهنتون</div>
            <div class="sidebar-subtitle">AI Student Assistant</div>
        </div>
        """
    )
    st.divider()
    st.html('<div class="example-title">💡 بېلګې پوښتنې</div>')

    for i, example in enumerate(EXAMPLES):
        if st.button(example, key=f"example_{i}", width="stretch"):
            st.session_state.pending = example

    st.write("")
    if st.button("🗑️ چت پاک کړئ", width="stretch"):
        st.session_state.history = []
        st.rerun()

    st.divider()
    st.html(
        """
        <div style="direction:rtl; text-align:right; color:#718096;
                    font-size:0.82rem; line-height:1.9;">
            <b>د مرستیال په اړه</b><br><br>
            دا سیستم د ننګرهار پوهنتون د زده کوونکو لپاره جوړ شوی ډیجیټل مرستیال دی.
            <br><br>
            د پوهنتون له معلوماتو څخه د پوښتنو ځوابولو لپاره مصنوعي ځیرکتیا کاروي.
        </div>
        """
    )

# =========================================================
# API KEY
# =========================================================
def secret_key():
    try:
        return st.secrets["GEMINI_API_KEY"]
    except Exception:
        return ""


api_key = os.environ.get("GEMINI_API_KEY") or secret_key() or st.session_state.get("api_key", "")

if not api_key:
    st.html(
        """
        <div class="hero"><div class="hero-content">
            <div class="hero-icon">🎓</div>
            <div class="hero-title">د ننګرهار پوهنتون د زده کوونکو ډیجیټل مرستیال</div>
            <div class="hero-subtitle">Nangarhar University · AI Student Support</div>
            <div class="status">🔐 API Key ته اړتیا ده</div>
        </div></div>
        """
    )
    st.html(
        """
        <div class="welcome">
            <div class="welcome-icon">🔑</div>
            <div class="welcome-title">Gemini API Key داخل کړئ</div>
            <div class="welcome-text">
                د چت بوټ د فعالولو لپاره خپل Gemini API Key لاندې داخل کړئ.<br>
                API Key به یوازې د دې Streamlit session لپاره وکارول شي.
            </div>
        </div>
        """
    )
    st.html(
        """
        <div class="api-box">
            <div class="api-title">🔐 Gemini API Key</div>
            <div class="api-description">
                خپل Gemini API Key دلته ولیکئ. له داخلولو وروسته به چت بوټ فعال شي.
            </div>
        </div>
        """
    )
    entered_api_key = st.text_input(
        "Gemini API Key",
        type="password",
        placeholder="Paste your Gemini API Key here...",
        label_visibility="collapsed",
    )
    if entered_api_key:
        st.session_state.api_key = entered_api_key.strip()
        st.rerun()

    st.info("💡 خپل Gemini API Key د Google AI Studio (aistudio.google.com/apikey) څخه ترلاسه کولی شئ.")
    st.stop()

# =========================================================
# READ THE QUESTION FIRST (so the welcome card can hide itself)
# =========================================================
typed = st.chat_input("خپله پوښتنه ولیکئ...")
question = typed or st.session_state.pop("pending", None)

# =========================================================
# MAIN HERO
# =========================================================
st.html(
    """
    <div class="hero"><div class="hero-content">
        <div class="hero-icon">🎓</div>
        <div class="hero-title">د ننګرهار پوهنتون د زده کوونکو ډیجیټل مرستیال</div>
        <div class="hero-subtitle">Nangarhar University · AI Student Support</div>
        <div class="status"><span class="status-dot"></span>AI Assistant Online</div>
    </div></div>
    """
)

# =========================================================
# DATA CHECK (owner only): open  your-link/?debug=1
# =========================================================
if st.query_params.get("debug") == "1":
    try:
        _c, _i, _v, rep = load_index(api_key, data_signature())
        lines = [f"**Text chunks:** {len(_c)}"]
        lines += [f"- ✅ {n}: {k} chunks" for n, k in rep["read"]]
        lines += [f"- ⚠️ {n}: {why}" for n, why in rep["skipped"]]
        lines.append(f"**Photos found in images/:** {', '.join(rep['photos']) or 'none'}")
        st.info("\n".join(lines))
    except Exception as e:
        st.error(f"Data check failed: {e}")

# =========================================================
# WELCOME (only before the first question)
# =========================================================
if not st.session_state.history and not question:
    st.html(
        """
        <div class="welcome">
            <div class="welcome-icon">🤖</div>
            <div class="welcome-title">ښه راغلاست! 👋</div>
            <div class="welcome-text">
                زه د ننګرهار پوهنتون د زده کوونکو ډیجیټل مرستیال یم.<br>
                زه کولی شم د موجودو پوهنتوني معلوماتو پر بنسټ ستاسو پوښتنو ته ځواب ووایم.
                <br><br>
                خپله پوښتنه په پښتو ولیکئ.
            </div>
        </div>
        """
    )

# =========================================================
# CHAT HISTORY
# =========================================================
for item in st.session_state.history:
    with st.chat_message(item["role"], avatar="🧑‍🎓" if item["role"] == "user" else "🎓"):
        st.write(item["text"])
        show_photos(item.get("images", []))

# =========================================================
# PROCESS QUESTION
# =========================================================
TYPING_HTML = '<div class="typing"><span></span><span></span><span></span></div>'

if question:
    if st.session_state.count >= MAX_QUESTIONS:
        st.warning("تاسو د دې ناستې د پوښتنو حد ته رسېدلي یاست. مهرباني وکړئ وروسته بیا هڅه وکړئ 🙏")
    else:
        st.session_state.count += 1
        past = list(st.session_state.history)  # earlier messages, without the new question
        st.session_state.history.append({"role": "user", "text": question, "images": []})
        with st.chat_message("user", avatar="🧑‍🎓"):
            st.write(question)

        # short follow-ups ("and the fees?") need the previous question to find the right documents
        last_user = next((m["text"] for m in reversed(past) if m["role"] == "user"), "")
        search_text = f"{last_user} {question}" if len(question) < 30 else question

        with st.chat_message("assistant", avatar="🎓"):
            box = st.empty()
            box.markdown(TYPING_HTML, unsafe_allow_html=True)  # animated dots while thinking
            reply, pics = "", []
            try:
                texts, pics = retrieve(search_text, api_key)
                context = "\n\n".join(texts) + photo_block(pics)
                for piece in stream_answer(question, context, api_key, past):
                    reply += piece
                    box.markdown(split_photos(reply, pics)[0] + " ▌")
            except Exception as e:
                reply = f"Error: {e}"

            ok = not reply.startswith("Error:")
            visible, shown = split_photos(reply, pics) if ok else (reply, [])
            if ok and len(visible) > 80:
                visible += "\n\n" + random.choice(CLOSINGS)  # friendly closing line
            box.markdown(visible)
            show_photos(shown)

        if ok:
            st.session_state.history.append({"role": "assistant", "text": visible, "images": shown})
        else:
            st.session_state.history.pop()  # do not keep a failed turn in the conversation

# =========================================================
# FOOTER
# =========================================================
st.html(
    """
    <div class="footer">
        🎓 د ننګرهار پوهنتون AI Student Assistant<br>
        د دقیقو او تازه معلوماتو لپاره د پوهنتون رسمي ویب پاڼه وګورئ.<br>
        <span style="font-family:Segoe UI;">nu.edu.af</span><br>
        مهرباني وکړئ شخصي معلومات مه لیکئ.
    </div>
    """
)